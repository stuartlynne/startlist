
import sys
import os
import requests


import openpyxl
from datetime import datetime
import openpyxl
import csv
from operator import itemgetter
from urllib.parse import urljoin

class Sheet:
    def __init__(self, worksheet, fieldnames=None, title=None):
        self.title = title
        self.worksheet = worksheet
        self.fieldnames = fieldnames
        self.row_idx = 1

        self.worksheet.title = title
        self.worksheet.freeze_panes = 'A2'
        self.headers = [k for k in fieldnames.keys()]
        self.fields = [v for v in fieldnames.values()]
        self.fieldCount = [0 for k in fieldnames.keys()]
        print('Sheet:', self.title, file=sys.stderr)
        print('Headers:', self.headers, file=sys.stderr)
        print('Fields:', self.fields, file=sys.stderr)

        self.worksheet.append(self.headers)
        header_font = openpyxl.styles.Font(bold=True)
        for col_idx in range(1, len(self.headers)+1):
            self.worksheet.cell(row=1, column=col_idx).font = header_font
            self.worksheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = self.fields[col_idx-1][1]
            print('self.fields:', self.fields[col_idx-1], file=sys.stderr)
            self.worksheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].alignment = openpyxl.styles.Alignment(horizontal=self.fields[col_idx-1][0], wrap_text=True)

        self.worksheet.auto_filter.ref = self.worksheet.dimensions
        self.worksheet.row_dimensions[1].height = 30

    def write_row(self, row=None, ):
        print('WriteRow:', row, file=sys.stderr)
        row[0] = self.row_idx
        for idx, v in enumerate(row):
            if not self.fields[idx][2]:
                self.fieldCount[idx] += 1
            elif row[idx] == 'None':
                self.fieldCount[idx] += 1
            elif row[idx] == 'Y':
                self.fieldCount[idx] += 1

            print('Row:', v, file=sys.stderr)

        self.worksheet.append(row)
        for cell_idx, cell in enumerate(self.worksheet[self.row_idx]):
            cell.alignment = openpyxl.styles.Alignment(horizontal=self.fields[cell_idx][0], wrap_text=True)
        self.row_idx += 1

    def write_totals(self, ):
        for field_idx in range(0, len(self.fieldCount)):
            if field_idx == 0:
                self.fieldCount[0] = 'Total'
            elif field_idx in [2, 3,]:
                self.fieldCount[field_idx] = ''
            else:
                self.fieldCount[field_idx] = f"{self.fieldCount[field_idx]} / {self.row_idx-self.fieldCount[field_idx]-1}"

        self.worksheet.append(self.fieldCount)
        for cell_idx, cell in enumerate(self.worksheet[self.row_idx]):
            cell.alignment = openpyxl.styles.Alignment(horizontal=self.fields[cell_idx][0], wrap_text=True)
        self.row_idx += 1
        self.worksheet.row_dimensions[self.row_idx].height = 30

class GenAudit:
    def __init__(self, racedb_host, date, competition_id, competition_name, ):
        print('GenAudit:', racedb_host, date, competition_id, competition_name, file=sys.stderr)
        self.racedb_host = racedb_host
        self.date = date
        self.competition_id = competition_id
        self.competition_name = competition_name
        self.events = []
        self.participants = []

    def add_event(self, event_id, event_name, event_start_time):
        self.events.append(event_id)

    def add_wave(self, event_id, wave_id, wave_name, start_offset, distance, laps, minutes, categories, ):
        pass

    def add_participant(self, event_id, wave_id, wave_name, participant_data):
        print('GenAudit.Participant:', wave_name, participant_data, file=sys.stderr)
        for k in ['last_name', 'first_name', 'bib', 'category_code']:
            participant_data[k] = str(participant_data[k]).title()
        for k in ['preregistered', 'tag_checked', 'license_checked', 'paid']:
            if participant_data[k]:
                participant_data[k] = 'Y'
            else:
                participant_data[k] = 'N'
        self.participants.append(participant_data)
        pass

    def save(self, ):
        #fieldnames = ['Last Name,First Name,Bib,Category,PreReg,tag,lic,paid']
        #dictnames = ['last_name', 'first_name', 'bib', 'category_code', 'preregistered', 'tag_checked', 
        #        'license_checked', 'paid']
        dictnames = {
                '#': '#',
                'last_name': 'Last Name',
                'first_name': 'First Name',
                'bib': 'Bib',
                'preregistered': 'Pre Reg',
                'tag_checked': 'Tag Read',
                'license_checked': 'Lic Chk',
                'paid': 'Paid',
                'category_code': 'Category',
                }
        fieldnames = {
            '#': ('center', 5, False, ),
            'Last Name': ('left', 20, False, ),
            'First Name': ('left', 20, False, ),
            'Bib': ('center', 7, True, ),
            'Pre Reg': ('center', 7, True,),
            'Tag Read': ('center', 7, True,),
            'Lic Chk': ('center', 7, True,),
            'Paid': ('center', 7, True,),
            'Category': ('left', 30, False,)
             }
        alignment = {'last_name': 'left',
                     'first_name': 'left',
                     'bib': 'center',
                     'preregistered': 'center',
                     'tag_checked': 'center',
                     'license_checked': 'center',
                     'paid': 'center',
                     'category_code': 'left',
                     }
         #align = [alignment[k] for k in dictnames.keys()]

        
        with open ('audit.csv', 'w') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=[k for k in fieldnames.keys()])
            writer.writeheader()
            for participant in sorted(self.participants, key=itemgetter('last_name', 'first_name')):
                print('Participant:', participant, file=sys.stderr)
                writer.writerow( {v: participant.get(k,'') for k,v in dictnames.items()})
                #print("Row: %s" % ( {v: participant[k] for k,v in dictnames.items()}), file=sys.stderr
        

        workbook = openpyxl.Workbook()
        particpantsWorksheet = Sheet(workbook.active, fieldnames=fieldnames, title="All Participants")
        nonregisteredWorksheet = Sheet(workbook.create_sheet(title="Non-Registered"), fieldnames=fieldnames, title="Non-Registered")


        #worksheet.title = "All Participants"
        #worksheet.freeze_panes = 'A2'
        #worksheet.append([v for v in dictnames.values()])
        #header_font = openpyxl.styles.Font(bold=True)
        #for col in range(1, len(dictnames)+1):
        #worksheet.cell(row=1, column=col).font = header_font



        for row_idx, participant in enumerate(sorted(self.participants, key=itemgetter('last_name', 'first_name')), start=2):
            row = [participant.get(k, '') for k in dictnames.keys()]
            print('Row:', row, file=sys.stderr)
            
            particpantsWorksheet.write_row(row=row, )

            if participant['preregistered'] == 'N':
                nonregisteredWorksheet.write_row(row=row, )

            #particpantsWorksheet.append(row)
            #for cell in particpantsWorksheet[row_idx]:
            #    if cell.column in [3, 4, 5, 6,]:
            #        cell.alignment = openpyxl.styles.Alignment(horizontal='center')
            #    #cell.alignment = openpyxl.styles.Alignment(horizontal='center')
                

        particpantsWorksheet.write_totals()
        nonregisteredWorksheet.write_totals()
        workbook.save('audit.xlsx')


