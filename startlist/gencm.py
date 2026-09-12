
import sys
import os
import requests


import openpyxl
from datetime import datetime
from urllib.parse import urljoin
from .filename import sanitize_filename_part

class GenCM:
    def __init__(self, racedb_host, date, competition_id, competition_name, series_all=False):
        print('GenCM:', racedb_host, date, competition_id, competition_name, file=sys.stderr)
        self.racedb_host = racedb_host
        self.date = date
        self.competition_id = competition_id
        self.competition_name = competition_name
        self.series_all = series_all
        self.events = {}

    def add_event(self, event_id, event_name, event_start_time):
        print(f"Adding event {event_id} {event_name} to competition", file=sys.stderr)
        self.event_name = event_name
        self.events[event_name] = event_id

    def add_wave(self, event_id, wave_id, wave_name, start_offset, distance, laps, minutes, categories, ):
        pass

    def add_participant(self, event_id, wave_id, wave_name, participant_data):
        pass

    def force_series_all(self, file_name):
        workbook = openpyxl.load_workbook(file_name)
        sheet_name = '--CrossMgr-Categories'
        if sheet_name not in workbook.sheetnames:
            print(f"Series all requested but {sheet_name} not found in {file_name}", file=sys.stderr)
            return

        worksheet = workbook[sheet_name]
        series_col = None
        for cell in worksheet[1]:
            if cell.value == 'Series':
                series_col = cell.column
                break

        if series_col is None:
            print(f"Series all requested but Series column not found in {file_name}", file=sys.stderr)
            return

        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=series_col, value=True)

        workbook.save(file_name)
        print(f"Forced Series TRUE in {file_name}", file=sys.stderr)

    def save(self, category_bib_ranges=None):
        #for i, event_id in enumerate(self.events, 1):
        for i, (event_name, event_id) in enumerate(self.events.items()):
            print(f"Downloading event {event_id}...", file=sys.stderr)
            url = urljoin(f"https://{self.racedb_host}/RaceDB/Competitions/CompetitionDashboard/{self.competition_id}/EventMassStartCrossMgr/", 
                          str(event_id))
            print('Download URL:', url, file=sys.stderr)
            response = requests.get(url)
            if response.status_code == 200:
                #file_name = f"{self.date}-{self.competition_name.replace(' ', '_')}_{i}.xlsx"
                #file_name = f"{self.date}-{event_name.replace(' ', '_')}.xlsx"
                #file_name = f"{self.date}-{self.competition_name.replace(' ', '_')}_{event_name.replace('/','_').replace(' ','_')}-r{i}.xlsx"
                competition_slug = sanitize_filename_part(self.competition_name, compact=True)
                event_slug = sanitize_filename_part(event_name)
                file_name = f"{self.date}-{competition_slug}-{event_slug}.xlsx"
                with open(file_name, "wb") as file:
                    file.write(response.content)
                if self.series_all:
                    self.force_series_all(file_name)
                print(f"Downloaded event {event_id} as {file_name}", file=sys.stdout)
            else:
                print(f"Failed to download event {self.event_id} from {url}", file=sys.stdout)
